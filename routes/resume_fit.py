"""Resume endpoints for HR AI Platform

Improvements:
- DOCX parsing via `python-docx`
- PDF text fallbacks: PyPDF2 -> pdfplumber -> OCR (pytesseract + OpenCV)
- Image OCR for uploaded images
- New `/resume/analyze` endpoint that returns skills and role predictions
"""
import os
from typing import List
import PyPDF2
from fastapi import APIRouter, File, UploadFile, HTTPException, Depends, Query, Form, Body, Request
from sqlalchemy.orm import Session
from db import models as db_models
from db.models import User, Resume, Job, JobMatch
from db.config import get_db_connection, close_db_connection
from db.database import get_db
from services import embedding_service, skill_extractor
from services.auth import get_current_user, get_current_user_optional

import io
import logging

logger = logging.getLogger(__name__)

# Module-level db reference for backward compat
class _db:
    @staticmethod
    def get_db_connection():
        return get_db_connection()
    @staticmethod
    def close_db_connection(conn):
        return close_db_connection(conn)
db = _db()

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

try:
    import pytesseract
except Exception:
    pytesseract = None

try:
    import cv2
except Exception:
    cv2 = None

try:
    import numpy as np
except Exception:
    np = None

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None

router = APIRouter(prefix="/resume", tags=["Resume"])
api_router = APIRouter(prefix="/api", tags=["API"])


def extract_text_from_path(path: str, filename: str) -> str:
    """Top-level helper for extracting text from different file types.

    Uses PyPDF2 -> pdfplumber -> OCR (pytesseract+OpenCV) for PDFs,
    python-docx for DOCX, and pytesseract+OpenCV for images.
    """
    text = ""
    ext = os.path.splitext(filename)[1].lower()

    # PDF: try PyPDF2, then pdfplumber, then OCR
    if ext == ".pdf":
        try:
            with open(path, "rb") as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                for page in pdf_reader.pages:
                    try:
                        page_text = page.extract_text() or ""
                        text += page_text + "\n"
                    except Exception:
                        continue
        except Exception:
            logger.exception("PyPDF2 failed to read PDF")

        if not text.strip() and pdfplumber is not None:
            try:
                with pdfplumber.open(path) as pdf:
                    for page in pdf.pages:
                        try:
                            p_text = page.extract_text() or ""
                            text += p_text + "\n"
                        except Exception:
                            continue
            except Exception:
                logger.exception("pdfplumber failed to read PDF")

        # If still empty, try OCR on each page. Prefer PyMuPDF (fitz) for
        # higher-quality rasterization; fall back to pdfplumber image rendering.
        if not text.strip() and pytesseract is not None and np is not None:
            DPI = int(os.getenv("PDF_OCR_DPI", "300"))

            # Try PyMuPDF first (high-quality rasterization)
            if fitz is not None:
                try:
                    doc = fitz.open(path)
                    for page in doc:
                        try:
                            mat = fitz.Matrix(DPI / 72.0, DPI / 72.0)
                            pix = page.get_pixmap(matrix=mat, alpha=False)
                            img_bytes = pix.tobytes("png")
                            # Convert bytes -> cv2 image if available
                            if cv2 is not None:
                                nparr = np.frombuffer(img_bytes, np.uint8)
                                img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
                                if img is None:
                                    continue
                                gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
                                gray = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
                                ocr_text = pytesseract.image_to_string(gray)
                            else:
                                # PIL fallback
                                from PIL import Image
                                from io import BytesIO
                                img_pil = Image.open(BytesIO(img_bytes)).convert('RGB')
                                ocr_text = pytesseract.image_to_string(img_pil)
                            text += ocr_text + "\n"
                        except Exception:
                            continue
                except Exception:
                    logger.exception("PyMuPDF rendering + OCR failed; falling back to pdfplumber")

            # Fallback to pdfplumber rendering with configurable DPI
            if not text.strip() and pdfplumber is not None and cv2 is not None:
                try:
                    with pdfplumber.open(path) as pdf:
                        for page in pdf.pages:
                            try:
                                pil_image = page.to_image(resolution=DPI).original
                                img = np.array(pil_image.convert('RGB'))
                                img = cv2.cvtColor(img, cv2.COLOR_RGB2BGR)
                                gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
                                gray = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
                                ocr_text = pytesseract.image_to_string(gray)
                                text += ocr_text + "\n"
                            except Exception:
                                continue
                except Exception:
                    logger.exception("pdfplumber OCR fallback failed")

    # DOCX parsing
    elif ext == ".docx":
        if DocxDocument is not None:
            try:
                doc = DocxDocument(path)
                paragraphs = [p.text for p in doc.paragraphs if p.text]
                text = "\n".join(paragraphs)
                # Extract tables
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text:
                                text += "\n" + cell.text
            except Exception:
                logger.exception("Failed to parse DOCX file")
        else:
            logger.warning("python-docx not available; cannot parse DOCX")

    # Images: OCR using pytesseract + OpenCV
    elif ext in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        if pytesseract is not None and cv2 is not None and np is not None:
            try:
                img = cv2.imread(path)
                if img is None:
                    # Try PIL fallback
                    from PIL import Image
                    img = np.array(Image.open(path).convert('RGB'))
                    img = cv2.cvtColor(img, cv2.COLOR_RGB2BGR)
                gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
                gray = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
                text = pytesseract.image_to_string(gray)
            except Exception:
                logger.exception("Image OCR failed")
        else:
            logger.warning("OCR dependencies missing; cannot OCR image")

    else:
        logger.warning("Unsupported file extension for text extraction: %s", ext)

    return text or ""

@router.post("/upload")
async def upload_resume(
    request: Request,
    file: UploadFile | None = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
    resume_text: str | None = Body(default=None),
):
    """Upload and process a resume from either a file or raw text."""
    if file is not None:
        allowed_types = [".pdf", ".doc", ".docx"]
        if not any(file.filename.endswith(ext) for ext in allowed_types):
            raise HTTPException(status_code=400, detail="Invalid file type")

        # Save file
        file_path = f"uploads/resumes/{current_user.id}_{file.filename}"
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, "wb") as buffer:
            buffer.write(await file.read())

        # Extract text using top-level helper
        resume_text = extract_text_from_path(file_path, file.filename)

    if resume_text is None and request.headers.get("content-type", "").startswith("application/json"):
        try:
            payload = await request.json()
            if isinstance(payload, dict):
                resume_text = payload.get("resume_text")
        except Exception:
            resume_text = None

    if not resume_text or not resume_text.strip():
        raise HTTPException(status_code=400, detail="Provide resume text or a valid resume file")

    # Detect language and extract skills (language-aware)
    try:
        from services.language_utils import detect_language
        lang = detect_language(resume_text)
    except Exception:
        lang = None

    skills = skill_extractor.extract_skills(resume_text, lang=lang)
    vector = embedding_service.embed_text(resume_text)

    # Convert vector to bytes hex safely
    try:
        import numpy as _np
        vec_np = _np.asarray(vector, dtype=_np.float32)
        vector_hex = vec_np.tobytes().hex()
    except Exception:
        import json
        vector_hex = json.dumps(vector.tolist() if hasattr(vector, "tolist") else vector)

    # Store in database
    resume_record = Resume(
        tenant_id=current_user.tenant_id,
        user_id=current_user.id,
        resume_text=resume_text,
        skills=skills,
        vector_embedding=vector_hex,
    )
    db.add(resume_record)
    db.commit()
    db.refresh(resume_record)

    return {
        'id': resume_record.id,
        'user_id': current_user.id,
        'skills': skills
    }


@router.post("/analyze")
async def analyze(
    file: UploadFile = File(None),
    resume_text: str = Form(None),
    job_description: str = Form(None),
    current_user: User | None = Depends(get_current_user_optional)
):
    """Analyze resume (file upload or raw text) and return skills + role prediction.

    - Accepts either an uploaded file or raw `resume_text` form field.
    - Optionally accepts `job_description` to compute match score.
    """
    if not file and not resume_text:
        raise HTTPException(status_code=400, detail="Provide a resume file or resume_text")

    if file:
        # Save and extract text. Guest users can still analyze a file without auth.
        user_prefix = current_user.id if current_user else "guest"
        file_path = f"uploads/resumes/{user_prefix}_{file.filename}"
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        with open(file_path, "wb") as buffer:
            buffer.write(await file.read())

        resume_text = extract_text_from_path(file_path, file.filename)

    # Detect language and extract skills (language-aware)
    try:
        from services.language_utils import detect_language
        lang = detect_language(resume_text)
    except Exception:
        lang = None

    skills = skill_extractor.extract_skills(resume_text, lang=lang)

    # Role prediction (baseline TF-IDF + LogisticRegression)
    role_scores = []
    try:
        from ai_engine.role_predictor import role_predictor
        role_scores = role_predictor.predict(resume_text, top_k=3)
    except Exception:
        logger.debug("Role predictor unavailable or failed (model not trained yet)")

    result = {
        'skills': skills,
        'role_predictions': role_scores,
        'resume_text_preview': (resume_text[:200] + '...') if resume_text and len(resume_text) > 200 else (resume_text or "")
    }

    if lang:
        result['language'] = lang

    if job_description:
        emb_resume = embedding_service.embed_text(resume_text)
        emb_job = embedding_service.embed_text(job_description)
        match_score = embedding_service.calculate_similarity(emb_resume, emb_job)
        result['match_score'] = float(match_score)

    return result


@router.post('/upload-voice')
async def upload_voice(
    file: UploadFile = File(...),
    language: str | None = Form(None),
    current_user: User | None = Depends(get_current_user_optional)
):
    """Upload an audio file (wav/mp3) and transcribe it to resume text, then analyze."""
    # Save audio file
    allowed = {'.wav', '.mp3', '.m4a', '.ogg', '.flac', '.webm', '.aac'}
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in allowed:
        raise HTTPException(status_code=400, detail='Unsupported audio format')
    user_prefix = current_user.id if current_user else "guest"
    file_path = f"uploads/resumes/{user_prefix}_audio_{file.filename}"
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    with open(file_path, 'wb') as f:
        f.write(await file.read())

    # Transcribe
    try:
        # local import to avoid circular imports
        from services.speech_to_text import transcribe_audio
        resume_text = transcribe_audio(file_path, language=language)
    except Exception as e:
        logger.exception('Speech-to-text failed')
        raise HTTPException(status_code=500, detail=str(e))

    # Reuse analyze logic: extract skills and role predictions
    try:
        from services.language_utils import detect_language
        detected = detect_language(resume_text)
    except Exception:
        detected = None

    skills = skill_extractor.extract_skills(resume_text, lang=detected)
    role_scores = []
    try:
        from ai_engine.role_predictor import role_predictor
        role_scores = role_predictor.predict(resume_text, top_k=3)
    except Exception:
        logger.debug('Role predictor unavailable (model not trained yet)')

    result = {
        'skills': skills,
        'role_predictions': role_scores,
        'resume_text': resume_text
    }
    if detected:
        result['language'] = detected
    return result


@router.post('/chat')
async def chat_resume(
    resume_id: int = Form(...),
    message: str = Form(...),
    current_user: User = Depends(get_current_user)
):
    """Simple chatbot against a stored resume. Returns extracted answer.

    If an LLM is configured the endpoint will try to use it for a richer
    response; otherwise a retrieval-based answer is returned.
    """
    # Fetch resume text from DB
    conn = db.get_db_connection()
    cursor = conn.cursor()
    query = "SELECT resume_text FROM resumes WHERE id = %s AND user_id = %s AND tenant_id = %s"
    cursor.execute(query, (resume_id, current_user.id, current_user.tenant_id))
    row = cursor.fetchone()
    cursor.close()
    db.close_db_connection(conn)
    if not row:
        raise HTTPException(status_code=404, detail='Resume not found')
    resume_text = row[0] or ""

    # Try LLM-assisted answer
    try:
        from services.llm import call_llm
        prompt = f"You are an assistant that answers questions about a candidate's resume.\nResume:\n{resume_text}\n\nQuestion: {message}\nAnswer concisely."
        ans = call_llm(prompt, max_tokens=256)
        return {'answer': ans}
    except NotImplementedError:
        # Gemini stub not implemented; fall through to retrieval
        pass
    except Exception:
        logger.debug('LLM call failed; falling back to retrieval', exc_info=True)

    # Retrieval-based fallback
    try:
        from services.qa_service import answer_from_resume
        ans = answer_from_resume(resume_text, message)
        return {'answer': ans}
    except Exception:
        logger.exception('QA failed')
        raise HTTPException(status_code=500, detail='Failed to answer')


@router.post('/rank')
async def rank_resumes(
    job_description: str = Form(...),
    top_k: int = Form(10),
    current_user: User = Depends(get_current_user)
):
    """Rank all resumes in DB by relevance to the given job description."""
    conn = db.get_db_connection()
    cursor = conn.cursor()
    query = "SELECT id, user_id, resume_text, skills, vector_embedding FROM resumes WHERE tenant_id = %s"
    cursor.execute(query, (current_user.tenant_id,))
    rows = cursor.fetchall()

    # Compute embedding for job
    emb_job = embedding_service.embed_text(job_description)

    results = []
    for r in rows:
        rid = r[0]
        resume_text = r[2] or ''
        # Compute similarity
        try:
            emb_resume = embedding_service.embed_text(resume_text)
            score = embedding_service.calculate_similarity(emb_resume, emb_job)
        except Exception:
            score = 0.0
        results.append({'resume_id': rid, 'user_id': r[1], 'score': float(score), 'skills': (r[3] or '').split(',')})

    # Sort and return top_k
    results.sort(key=lambda x: x['score'], reverse=True)
    return {'ranked': results[:top_k]}


@router.post('/suggest')
async def suggest_improvements(
    resume_id: int | None = Form(None),
    resume_text: str | None = Form(None),
    job_description: str | None = Form(None),
    current_user: User | None = Depends(get_current_user_optional)
):
    """Return automated improvement suggestions for a resume.

    Accepts either `resume_id` (existing resume) or raw `resume_text`.
    Optionally accepts `job_description` to tailor suggestions.
    """
    text = resume_text
    if resume_id and not current_user:
        raise HTTPException(status_code=401, detail='Authentication required to access saved resumes')

    if resume_id and not text:
        # fetch from DB
        conn = db.get_db_connection()
        cursor = conn.cursor()
        query = "SELECT resume_text FROM resumes WHERE id = %s AND user_id = %s AND tenant_id = %s"
        cursor.execute(query, (resume_id, current_user.id, current_user.tenant_id))
        row = cursor.fetchone()
        cursor.close()
        db.close_db_connection(conn)
        if not row:
            raise HTTPException(status_code=404, detail='Resume not found')
        text = row[0] or ''

    if not text:
        raise HTTPException(status_code=400, detail='Provide resume_id or resume_text')

    try:
        from services.resume_advisor import suggest_improvements
        suggestions = suggest_improvements(text, job_description)
        return {'suggestions': suggestions}
    except Exception:
        logger.exception('Failed to generate suggestions')
        raise HTTPException(status_code=500, detail='Failed to generate suggestions')

@api_router.api_route("/resumes/upload", methods=["GET", "POST"])
async def upload_resume_api(
    request: Request,
    payload: dict | None = Body(None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    resume_text = payload.get("resume_text") if payload else None
    return await upload_resume(request=request, file=None, resume_text=resume_text, current_user=current_user, db=db)


@api_router.get("/resumes")
async def get_my_resumes_api(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    return await get_my_resumes(current_user, db)


@api_router.api_route("/jobs/create", methods=["GET", "POST"])
async def create_job_api(
    payload: dict | None = Body(None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if not payload:
        raise HTTPException(status_code=400, detail="Job payload is required")

    job = Job(
        tenant_id=current_user.tenant_id,
        title=payload.get("title", ""),
        description=payload.get("description", ""),
        required_skills=payload.get("required_skills", []),
        department=payload.get("department", "")
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    return {
        "id": job.id,
        "title": job.title,
        "description": job.description,
        "required_skills": job.required_skills or [],
        "department": job.department,
    }


@api_router.api_route("/resumes/{resume_id}/match-jobs", methods=["GET", "POST"])
async def match_jobs_api(
    resume_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    return await match_resume(resume_id=resume_id, top_k=3, current_user=current_user, db=db)


@router.get("/my-resumes", response_model=List[dict])
async def get_my_resumes(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Get all resumes for current user"""
    resumes = db.query(Resume).filter(Resume.user_id == current_user.id).all()
    return [{
        'id': resume.id,
        'user_id': resume.user_id,
        'skills': resume.skills or []
    } for resume in resumes]

@router.post("/{resume_id}/match")
async def match_resume(
    resume_id: int,
    top_k: int = Query(default=3, gt=0),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Match a resume with available jobs"""
    resume = db.query(Resume).filter(Resume.id == resume_id, Resume.user_id == current_user.id, Resume.tenant_id == current_user.tenant_id).first()
    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    jobs = db.query(Job).filter(Job.tenant_id == current_user.tenant_id).all()
    matches = []
    for job in jobs:
        resume_embedding = embedding_service.embed_text(resume.resume_text or "")
        job_embedding = embedding_service.embed_text(job.description or "")
        similarity = embedding_service.calculate_similarity(resume_embedding, job_embedding)

        resume_skills = set(resume.skills or [])
        job_skills = set(job.required_skills or [])
        skill_overlap = len(resume_skills.intersection(job_skills))
        skill_overlap_score = skill_overlap / len(job_skills) if job_skills else 0

        matched_skills = list(resume_skills.intersection(job_skills))
        db.add(JobMatch(
            tenant_id=current_user.tenant_id,
            job_id=job.id,
            resume_id=resume_id,
            match_score=float(similarity),
            skills_matched=matched_skills,
        ))

        matches.append({
            "job": {
                "id": job.id,
                "title": job.title,
                "description": job.description,
                "required_skills": list(job_skills)
            },
            "match_score": float(similarity),
            "skill_overlap_score": float(skill_overlap_score),
            "matching_skills": matched_skills,
            "missing_skills": list(job_skills - resume_skills)
        })

    db.commit()

    matches.sort(key=lambda x: x["match_score"], reverse=True)
    return matches[:top_k]

@router.delete("/{resume_id}")
async def delete_resume(resume_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Delete a resume"""
    resume = db.query(Resume).filter(Resume.id == resume_id, Resume.user_id == current_user.id).first()
    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found")

    db.delete(resume)
    db.commit()

    file_path = f"uploads/resumes/{current_user.id}_{resume_id}"
    if os.path.exists(file_path):
        os.remove(file_path)

    return {"message": "Resume deleted successfully"}


# Batch resume upload
@router.post("/batch-upload")
async def batch_upload_resumes(
    files: List[UploadFile] = File(...),
    job_description: str = Form(None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Upload multiple resumes at once for batch processing.
    
    Returns a task ID for tracking the batch processing status.
    """
    from services.tasks.resume_tasks import batch_process_resumes
    
    # Save files temporarily
    file_infos = []
    for file in files:
        allowed_types = [".pdf", ".doc", ".docx"]
        if not any(file.filename.endswith(ext) for ext in allowed_types):
            raise HTTPException(status_code=400, detail=f"Invalid file type: {file.filename}")
        
        file_path = f"uploads/resumes/{current_user.id}_{file.filename}"
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        
        with open(file_path, "wb") as buffer:
            buffer.write(await file.read())
        
        file_infos.append({
            "file_path": file_path,
            "filename": file.filename,
            "job_description": job_description
        })
    
    # Queue batch processing task
    task = batch_process_resumes.delay(file_infos, current_user.id)
    
    return {
        "task_id": task.id,
        "status": "queued",
        "total_files": len(file_infos),
        "message": "Batch processing started. Use task_id to check status."
    }


# Get batch processing status
@router.get("/batch-status/{task_id}")
async def get_batch_status(task_id: str, current_user: User = Depends(get_current_user)):
    """Get the status of a batch processing task."""
    from services.celery_app import celery_app
    
    result = celery_app.AsyncResult(task_id)
    
    return {
        "task_id": task_id,
        "status": result.status,
        "result": result.result if result.ready() else None
    }


# Rank all resumes for a job description
@router.post("/rank-all")
async def rank_all_resumes(
    job_description: str = Form(...),
    top_k: int = Form(10),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Rank all resumes in the database against a job description.
    
    Returns top_k matches with scores.
    """
    resumes = db.query(Resume).all()
    
    if not resumes:
        return {"ranked": [], "message": "No resumes in database"}
    
    job_embedding = embedding_service.embed_text(job_description)
    
    results = []
    for resume in resumes:
        try:
            resume_embedding = embedding_service.embed_text(resume.resume_text or "")
            score = embedding_service.calculate_similarity(resume_embedding, job_embedding)
        except Exception:
            score = 0.0
        
        results.append({
            "resume_id": resume.id,
            "user_id": resume.user_id,
            "score": float(score),
            "skills": resume.skills or []
        })
    
    results.sort(key=lambda x: x["score"], reverse=True)
    
    return {"ranked": results[:top_k]}


# Custom skill dictionary upload
@router.post("/skill-dictionary")
async def upload_skill_dictionary(
    file: UploadFile = File(...),
    language: str = Form("en"),
    current_user: User = Depends(get_current_user)
):
    """
    Upload a custom skill dictionary CSV for skill extraction.
    
    Expected CSV format: skill (one per row)
    Optional columns: category, aliases
    """
    import csv
    import io
    
    if not file.filename.endswith(".csv"):
        raise HTTPException(status_code=400, detail="File must be CSV format")
    
    content = await file.read()
    text = content.decode("utf-8")
    
    reader = csv.DictReader(io.StringIO(text))
    
    skills = []
    for row in reader:
        skill = row.get("skill") or row.get("Skill") or row.get("SKILL")
        if skill:
            skills.append({
                "skill": skill.strip(),
                "category": row.get("category", "General"),
                "aliases": [a.strip() for a in row.get("aliases", "").split(";") if a.strip()]
            })
    
    # Save to language-specific file
    lang_dir = "data"
    os.makedirs(lang_dir, exist_ok=True)
    filepath = os.path.join(lang_dir, f"skills_{language}.csv")
    
    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["skill"])
        for s in skills:
            writer.writerow([s["skill"]])
    
    # Reload skill extractor
    from services.skill_extractor import _SKILLS_BY_LANG, _load_skill_list
    _SKILLS_BY_LANG.pop(language, None)
    _load_skill_list(language)
    
    return {
        "status": "success",
        "skills_loaded": len(skills),
        "language": language,
        "file": filepath
    }


# Export resume to HRIS format
@router.get("/{resume_id}/export")
async def export_resume(
    resume_id: int,
    format: str = Query("json", regex="^(json|workday|bamboohr)$"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Export resume in HRIS-compatible format.
    
    Supported formats:
    - json: Generic JSON format
    - workday: Workday-compatible format
    - bamboohr: BambooHR-compatible format
    """
    resume = db.query(Resume).filter(Resume.id == resume_id, Resume.user_id == current_user.id).first()
    
    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found")
    
    base_data = {
        "resume_id": resume.id,
        "user_id": resume.user_id,
        "skills": resume.skills or [],
        "created_at": resume.created_at.isoformat() if resume.created_at else None,
        "text_preview": (resume.resume_text or "")[:500]
    }
    
    if format == "json":
        return base_data
    
    elif format == "workday":
        return {
            "Worker_Data": {
                "Worker_ID": str(resume.user_id),
                "Resume": {
                    "Resume_ID": str(resume.id),
                    "Skills": [{"Skill_Name": s} for s in (resume.skills or [])],
                    "Resume_Text": resume.resume_text or "",
                    "Created_Date": resume.created_at.isoformat() if resume.created_at else None
                }
            }
        }
    
    elif format == "bamboohr":
        return {
            "employeeId": resume.user_id,
            "resume": {
                "id": resume.id,
                "skills": resume.skills or [],
                "text": resume.resume_text or "",
                "dateCreated": resume.created_at.isoformat() if resume.created_at else None
            }
        }


# Get skill dictionary
@router.get("/skill-dictionary/{language}")
async def get_skill_dictionary(language: str = "en"):
    """Get the current skill dictionary for a language."""
    from services.skill_extractor import _load_skill_list
    
    skills = _load_skill_list(language)
    
    return {
        "language": language,
        "count": len(skills),
        "skills": skills
    }
